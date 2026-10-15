"""
Seeds business_taxonomy, attribute_taxonomy, and signal_taxonomy
as GLOBAL defaults (tenant_id = NULL) from the L1/L2 source workbook.
Idempotent: safe to re-run, skips rows that already exist.
"""
import os
import openpyxl
from supabase import create_client
from dotenv import load_dotenv

load_dotenv()

SUPABASE_URL = os.environ["SUPABASE_URL"]
SUPABASE_SERVICE_ROLE_KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
XLSX_PATH = "iThink - L1 & L2 Details.xlsx"  # place this file in your repo root

# service_role key bypasses RLS entirely -- required here because we're
# inserting tenant_id = NULL rows, which no tenant-scoped policy allows.
# NEVER use this key in frontend/client code, only in trusted backend scripts.
sb = create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]

    vertical = []  # (l1_name, l2_name, applicable_product_type)
    for row in ws.iter_rows(min_row=13, max_row=181, min_col=1, max_col=4, values_only=True):
        if row[0] is not None and row[1] != "L1 Name":
            vertical.append((row[1], row[2], row[3]))

    horizontal = []  # (l1_name, l2_name)
    for row in ws.iter_rows(min_row=13, max_row=181, min_col=6, max_col=8, values_only=True):
        if row[0] is not None and row[1] != "Horizontal Layer L1":
            horizontal.append((row[1], row[2]))

    return vertical, horizontal


def get_or_create_business_l1(l1_name, cache):
    if l1_name in cache:
        return cache[l1_name]
    existing = sb.table("business_taxonomy").select("id").is_("tenant_id", "null").eq("l1_name", l1_name).execute()
    if existing.data:
        cache[l1_name] = existing.data[0]["id"]
    else:
        inserted = sb.table("business_taxonomy").insert({"l1_name": l1_name, "tenant_id": None}).execute()
        cache[l1_name] = inserted.data[0]["id"]
    return cache[l1_name]


def seed_attribute_taxonomy(vertical_rows):
    l1_cache = {}
    inserted, skipped = 0, 0
    for l1_name, l2_name, product_type in vertical_rows:
        business_l1_id = get_or_create_business_l1(l1_name, l1_cache)
        existing = (
            sb.table("attribute_taxonomy")
            .select("id")
            .is_("tenant_id", "null")
            .eq("business_l1_id", business_l1_id)
            .eq("l2_name", l2_name)
            .execute()
        )
        if existing.data:
            skipped += 1
            continue
        sb.table("attribute_taxonomy").insert({
            "tenant_id": None,
            "business_l1_id": business_l1_id,
            "l2_name": l2_name,
            "applicable_product_type": product_type,
        }).execute()
        inserted += 1
    print(f"attribute_taxonomy: {inserted} inserted, {skipped} already existed")


def seed_signal_taxonomy(horizontal_rows):
    # Add Severity and Importance manually -- confirmed additions not in the source sheet
    extra = [
        ("Severity", "Low"), ("Severity", "Medium"), ("Severity", "High"),
        ("Importance", "Low"), ("Importance", "Medium"), ("Importance", "High"),
    ]
    all_rows = horizontal_rows + extra

    inserted, skipped = 0, 0
    for l1_name, l2_name in all_rows:
        existing = (
            sb.table("signal_taxonomy")
            .select("id")
            .is_("tenant_id", "null")
            .eq("l1_name", l1_name)
            .eq("l2_name", l2_name)
            .execute()
        )
        if existing.data:
            skipped += 1
            continue
        sb.table("signal_taxonomy").insert({
            "tenant_id": None,
            "l1_name": l1_name,
            "l2_name": l2_name,
        }).execute()
        inserted += 1
    print(f"signal_taxonomy: {inserted} inserted, {skipped} already existed")


if __name__ == "__main__":
    vertical_rows, horizontal_rows = parse_workbook(XLSX_PATH)
    print(f"Parsed {len(vertical_rows)} vertical attributes, {len(horizontal_rows)} horizontal signals")
    seed_attribute_taxonomy(vertical_rows)
    seed_signal_taxonomy(horizontal_rows)